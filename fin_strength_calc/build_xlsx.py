#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""fin_thickness_calc.py 의 CLI 계산기를 라이브 엑셀 수식 워크북으로 변환.

매크로 없음(.xlsx). 모든 결과는 입력 셀을 참조하는 네이티브 수식으로 계산되어,
차압/온도/피치를 바꾸면 자동 재계산된다. 좌굴 최소두께(반복해)는 후보 두께 스캔
테이블 + MINIFS 로 대신한다.

셀 맵 (계산기 시트, 값은 C열):
  입력   C5 T · C6 ΔP_t · C7 ΔP_c · C8 p · C9 p_adj · C10 h · C11 K ·
         C12 f(스테이) · C13 공차 · C14 w · C15 P_min · C16 t_sel
  상수   C18 ν · C19 DF
  물성   C21 S(T) · C22 Sy(T) · C23 E(T) · C24 온도유효성
  결과   C26 인장 · C27 좌굴 · C28 요구(유효) · C29 요구(공칭) ·
         C30 두께판정 · C31 피치판정 · C32 높이판정
  파팅시트 C34 p̄ · C35 t_fin_eff · C36 스팬 · C37 M_d · C38 t_ps · C39 판정
  검증   C41 t_eff · C42 fσ_t · C43 인장판정 · C44 σ_c · C45 σ_cr ·
         C46 허용 · C47 압축판정 · C48 종합   (G45 = σ_cr,E 헬퍼)
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter

OUT = "/home/user/PFHE/fin_strength_calc/PFHE_fin_calc.xlsx"
CG = "계산기"

S_TABLE = [(40, 22.8), (65, 22.8), (93, 22.6), (121, 21.4),
           (149, 18.6), (177, 14.5), (204, 11.0)]
SY_TABLE = [(40, 34.5), (93, 33.1), (149, 29.6), (204, 24.1)]
E_TABLE = [(25, 68900), (93, 66200), (149, 63400), (204, 60000)]
SCAN_LO, SCAN_HI, SCAN_STEP = 0.010, 1.000, 0.005

ARIAL = "Arial"
f_title = Font(name=ARIAL, size=14, bold=True, color="FFFFFF")
f_sub = Font(name=ARIAL, size=9, italic=True, color="F2F2F2")
f_hdr = Font(name=ARIAL, size=11, bold=True, color="FFFFFF")
f_lbl = Font(name=ARIAL, size=10)
f_lbl_b = Font(name=ARIAL, size=10, bold=True)
f_in = Font(name=ARIAL, size=10, bold=True, color="0000FF")
f_out = Font(name=ARIAL, size=10)
f_out_b = Font(name=ARIAL, size=10, bold=True)
f_note = Font(name=ARIAL, size=8, italic=True, color="555555")
f_warn = Font(name=ARIAL, size=10, bold=True, color="C00000")
fill_title = PatternFill("solid", fgColor="1F4E78")
fill_hdr = PatternFill("solid", fgColor="2E75B6")
fill_in = PatternFill("solid", fgColor="FFF2CC")
fill_res = PatternFill("solid", fgColor="E2EFDA")
fill_prop = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
ctr = Alignment(horizontal="center", vertical="center")
left = Alignment(horizontal="left", vertical="center", wrap_text=True)
right = Alignment(horizontal="right", vertical="center")

wb = openpyxl.Workbook()

# ---------------- 물성DB ----------------
ws_m = wb.active
ws_m.title = "물성DB"
ws_m.sheet_view.showGridLines = False
ws_m["A1"] = "재료 물성 데이터 — AL 3003-O (ASME Sec. II-D 대표값)"
ws_m["A1"].font = Font(name=ARIAL, size=12, bold=True)
ws_m["A2"] = "⚠️ 대표값 — 실제 설계 전 적용 코드 연도판 값으로 교체할 것. 온도 오름차순 유지."
ws_m["A2"].font = f_warn

def write_table(ws, sr, sc, title, pairs, unit):
    c0, c1 = get_column_letter(sc), get_column_letter(sc + 1)
    ws.cell(sr, sc, title).font = f_lbl_b
    for k, txt in enumerate(("온도(°C)", unit)):
        c = ws.cell(sr + 1, sc + k, txt)
        c.font = f_hdr; c.fill = fill_hdr; c.alignment = ctr
    for i, (t, v) in enumerate(pairs):
        r = sr + 2 + i
        for k, val in enumerate((t, v)):
            c = ws.cell(r, sc + k, val)
            c.font = f_out; c.alignment = ctr; c.border = box
    first, lastr = sr + 2, sr + 1 + len(pairs)
    return f"물성DB!${c0}${first}:${c0}${lastr}", f"물성DB!${c1}${first}:${c1}${lastr}"

T_S, V_S = write_table(ws_m, 4, 1, "허용응력 S(T)", S_TABLE, "S (MPa)")
T_SY, V_SY = write_table(ws_m, 4, 4, "항복강도 Sy(T)", SY_TABLE, "Sy (MPa)")
T_E, V_E = write_table(ws_m, 4, 7, "탄성계수 E(T)", E_TABLE, "E (MPa)")
for col in "ABDEGH":
    ws_m.column_dimensions[col].width = 11
for name, ref in [("T_S", T_S), ("V_S", V_S), ("T_SY", T_SY), ("V_SY", V_SY),
                  ("T_E", T_E), ("V_E", V_E)]:
    wb.defined_names.add(DefinedName(name, attr_text=ref))

def interp(tc, TT, VV):
    m = f"MATCH({tc},{TT},1)"
    return (f"IF({tc}<=MIN({TT}),INDEX({VV},1),"
            f"IF({tc}>=MAX({TT}),INDEX({VV},COUNT({TT})),"
            f"INDEX({VV},{m})+({tc}-INDEX({TT},{m}))"
            f"/(INDEX({TT},{m}+1)-INDEX({TT},{m}))"
            f"*(INDEX({VV},{m}+1)-INDEX({VV},{m}))))")

# ---------------- 좌굴스캔 ----------------
ws_s = wb.create_sheet("좌굴스캔")
ws_s.sheet_view.showGridLines = False
for j, h in enumerate(["t 후보(mm)", "σ_c(MPa)", "σ_cr,E(MPa)",
                       "σ_cr Johnson(MPa)", "허용(MPa)", "OK(1/0)"], 1):
    c = ws_s.cell(1, j, h); c.font = f_hdr; c.fill = fill_hdr; c.alignment = ctr
ws_s["H1"] = "핀 좌굴 최소두께 = ΔP_c>0 일 때 OK=1 인 최소 t 후보 (계산기!C27, MINIFS)"
ws_s["H1"].font = f_note
n = int(round((SCAN_HI - SCAN_LO) / SCAN_STEP)) + 1
for i in range(n):
    r = i + 2
    ws_s.cell(r, 1, round(SCAN_LO + i * SCAN_STEP, 4)).number_format = "0.000"
    ws_s.cell(r, 2, f"={CG}!$C$7*{CG}!$C$8/A{r}")
    ws_s.cell(r, 3, f"=PI()^2*{CG}!$C$23/(12*(1-{CG}!$C$18^2))"
                    f"*(A{r}/({CG}!$C$11*{CG}!$C$10))^2")
    ws_s.cell(r, 4, f"=IF(C{r}<={CG}!$C$22/2,C{r},{CG}!$C$22*(1-{CG}!$C$22/(4*C{r})))")
    ws_s.cell(r, 5, f"=MIN(D{r}/{CG}!$C$19,{CG}!$C$21)")
    ws_s.cell(r, 6, f"=IF(B{r}<=E{r},1,0)")
    for j in range(2, 6):
        ws_s.cell(r, j).number_format = "0.00"
lastr = n + 1
wb.defined_names.add(DefinedName("BucklingT", attr_text=f"좌굴스캔!$A$2:$A${lastr}"))
wb.defined_names.add(DefinedName("BucklingOK", attr_text=f"좌굴스캔!$F$2:$F${lastr}"))
for col, w in zip("ABCDEF", (11, 11, 13, 17, 11, 9)):
    ws_s.column_dimensions[col].width = w

# ---------------- 계산기 ----------------
ws = wb.create_sheet("계산기")
wb.move_sheet("계산기", -(len(wb.sheetnames) - 1))
ws.sheet_view.showGridLines = False
for col, w in zip("ABCDE", (2, 31, 15, 8, 48)):
    ws.column_dimensions[col].width = w

def put(row, label, value=None, unit=None, note=None, *, kind="", num="0.000"):
    b = ws.cell(row, 2, label); b.font = f_lbl; b.alignment = left
    c = ws.cell(row, 3)
    if value is not None:
        c.value = value
    c.number_format = num; c.alignment = right; c.border = box
    styles = {"in": (f_in, fill_in), "res": (f_out_b, fill_res),
              "prop": (f_out, fill_prop)}
    if kind in styles:
        c.font, c.fill = styles[kind]
    else:
        c.font = f_out
    if unit:
        u = ws.cell(row, 4, unit); u.font = f_note; u.alignment = ctr
    if note:
        nn = ws.cell(row, 5, note); nn.font = f_note; nn.alignment = left
    return c

def section(row, text):
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
    c = ws.cell(row, 2, text); c.font = f_hdr; c.fill = fill_hdr
    c.alignment = Alignment(horizontal="left", vertical="center")

ws.merge_cells("B1:E1")
ws["B1"] = "PFHE Plain 핀 적정두께 계산기 (ALPEMA 2024 기반 유도)"
ws["B1"].font = f_title; ws["B1"].fill = fill_title
ws["B1"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 24
ws.merge_cells("B2:E2")
ws["B2"] = ("예비설계용 · 단위 MPa/mm · 파랑=입력 셀 · 인증설계는 "
            "ALPEMA 5.15.1.1(승인 계산법)/5.15.1.2(파열시험)에 따름")
ws["B2"].font = Font(name=ARIAL, size=9, italic=True, color="333333")
ws["B2"].alignment = Alignment(horizontal="center")
ws.merge_cells("B3:E3")
ws["B3"] = ('=IF($C$5>204,"⚠️ 설계온도 204°C 초과 — 3003 재질한계(ALPEMA Table 6-1) '
            '벗어남. 결과 무효!","")')
ws["B3"].font = f_warn; ws["B3"].alignment = Alignment(horizontal="center")

section(4, "① 입력 (파랑 셀을 편집)")
put(5, "설계온도 T", 65, "°C", "≤ 204 (3003 재질한계)", kind="in", num="0.0")
put(6, "핀 인장 유발 차압 ΔP_t", 5.0, "MPa", "통상 = 층 설계압력(MAWP)", kind="in")
put(7, "핀 압축 유발 차압 ΔP_c", 0.6, "MPa", "인접층 고압/자층 감압 시 (없으면 0)", kind="in")
put(8, "자층 핀 피치 p", 1.4, "mm", "ALPEMA 1.0–4.5", kind="in")
put(9, "인접층 핀 피치 p_adj", 0.0, "mm", "0=자층과 동일. 파팅시트 스팬용", kind="in")
put(10, "핀 높이 h", 6.5, "mm", "ALPEMA 2.0–12", kind="in", num="0.0")
put(11, "좌굴 유효길이계수 K", 0.7, "-", "기본 0.7, 구속 불확실 시 1.0", kind="in", num="0.00")
put(12, "인장 스테이 계수 f", 1.0, "-", "UG-50(b) 준용 시 1.10", kind="in", num="0.00")
put(13, "제작 공차계수", 1.0, "-", "0<f≤1 (예 0.9=공칭의 90%만 유효)", kind="in", num="0.00")
put(14, "파팅시트 차압 w=|P_A−P_B|", 5.0, "MPa", "층간 순차압", kind="in")
put(15, "균형가압 P_min", 5.0, "MPa", "양면 동시가압 낮은쪽(핀 어긋남 항)", kind="in")
put(16, "선정 핀 두께(공칭) t_sel", 0.30, "mm", "검증용. 0이면 검증 생략", kind="in")

section(17, "② 상수 및 온도별 물성 (자동)")
put(18, "포아송비 ν", 0.33, "-", "알루미늄", kind="prop", num="0.00")
put(19, "좌굴 설계계수 DF", 3.0, "-", "ASME UG-28 준용", kind="prop", num="0.0")
put(21, "허용응력 S(T)", f"={interp('$C$5', 'T_S', 'V_S')}", "MPa",
    "3003-O, 온도 선형보간", kind="prop", num="0.00")
put(22, "항복강도 Sy(T)", f"={interp('$C$5', 'T_SY', 'V_SY')}", "MPa",
    "Johnson 좌굴 천이용", kind="prop", num="0.00")
put(23, "탄성계수 E(T)", f"={interp('$C$5', 'T_E', 'V_E')}", "MPa",
    "좌굴 임계응력용", kind="prop", num="0")
put(24, "온도 유효성", '=IF($C$5>204,"온도초과(>204°C)","OK")', "-",
    "204°C 초과 시 결과 무효", kind="prop", num="General")

section(25, "③ 결과 — 핀 최소두께")
put(26, "[1] 핀 인장 최소두께 (식2)", "=IF($C$6<=0,0,$C$12*$C$6*$C$8/($C$21+$C$12*$C$6))",
    "mm", "t ≥ f·ΔP_t·p/(S+f·ΔP_t)", kind="res")
put(27, "[2] 핀 좌굴 최소두께 (식3~5)",
    "=IF($C$7<=0,0,IF(COUNTIF(BucklingOK,1)=0,999,_xlfn.MINIFS(BucklingT,BucklingOK,1)))",
    "mm", "Johnson 천이·전피치 하중폭 (좌굴스캔 시트)", kind="res")
ws.cell(27, 5).value = ('=IF(AND($C$7>0,COUNTIF(BucklingOK,1)=0),'
                        '"⚠️ 좌굴 소요두께 1.0mm 초과 — 표준 핀 불가",'
                        '"Johnson 천이·전피치 하중폭 (좌굴스캔 시트)")')
put(28, "[3] 요구 핀두께 (유효)", "=MAX($C$26,$C$27)", "mm",
    "max(인장, 좌굴)", kind="res")
put(29, "    요구 핀두께 (공칭)", '=IF($C$5>204,"온도초과",$C$28/$C$13)', "mm",
    "유효 / 공차계수", kind="res")
put(30, "  ALPEMA 두께범위 판정",
    '=IF($C$5>204,"—",IF(AND($C$29>=0.15,$C$29<=0.7),"OK 범위내","⚠️ 범위밖(0.15-0.7)"))',
    "-", "제작 가능 범위 0.15–0.7 mm", kind="prop", num="General")
put(31, "  핀 피치 판정",
    '=IF(AND($C$8>=1,$C$8<=4.5),"OK 범위내","⚠️ 범위밖(1.0-4.5)")', "-",
    "ALPEMA 5.11.4", kind="prop", num="General")
put(32, "  핀 높이 판정",
    '=IF(AND($C$10>=2,$C$10<=12),"OK 범위내","⚠️ 범위밖(2-12)")', "-",
    "ALPEMA 5.11.4", kind="prop", num="General")

section(33, "④ 파팅시트 두께 (식6~7)")
put(34, "지배 피치 p̄", "=IF($C$9>0,MAX($C$8,$C$9),$C$8)", "mm",
    "max(자층, 인접층 피치)", kind="prop")
put(35, "스팬 공제 핀두께 t_fin,eff",
    "=IF($C$9>$C$8,0,IF($C$16>0,$C$16*$C$13,$C$28))", "mm",
    "인접층 지배 시 0(보수) — 문서 §6.1", kind="prop")
put(36, "스팬 s", "=$C$34-$C$35", "mm", "s = p̄ − t_fin,eff", kind="prop")
put(37, "설계 모멘트 M_d", "=$C$14*$C$36^2/12+$C$15*($C$34-$C$35)*$C$34/8",
    "N·mm/mm", "w·s²/12 + P_min·(p̄−t)·p̄/8", kind="prop", num="0.0000")
put(38, "파팅시트 최소두께 t_ps",
    '=IF($C$36<=0,"스팬오류",SQRT(4*$C$37/$C$21))', "mm",
    "σ_b ≤ 1.5·S", kind="res")
put(39, "  ALPEMA 시트범위 판정",
    '=IF(NOT(ISNUMBER($C$38)),"—",IF($C$38>2,"⚠️ 범위밖(>2.0)",'
    'IF($C$38<0.8,"소요<0.8 → 표준 0.8 적용","OK 범위내")))', "-",
    "ALPEMA 5.11.1: 0.8–2.0 mm", kind="prop", num="General")

section(40, "⑤ 선정 두께 검증 (t_sel 입력 시)")
put(41, "유효 핀두께 t_eff", '=IF($C$16>0,$C$16*$C$13,"")', "mm",
    "공칭 × 공차계수", kind="prop")
put(42, "인장 f·σ_t",
    '=IF($C$16>0,IF($C$16*$C$13>=$C$8,"두께≥피치",'
    'IF($C$6>0,$C$12*$C$6*($C$8-$C$41)/$C$41,0)),"")', "MPa",
    "vs S(T)", kind="prop", num="0.00")
put(43, "  인장 판정",
    '=IF($C$16>0,IF(NOT(ISNUMBER($C$42)),"—",'
    'IF($C$42<=$C$21,"합격 ✓","불합격 ✗")),"")', "-", "", kind="res",
    num="General")
put(44, "압축 σ_c",
    '=IF(AND($C$16>0,$C$7>0),IF($C$16*$C$13>=$C$8,"두께≥피치",'
    '$C$7*$C$8/$C$41),"")', "MPa", "ΔP_c·p/t_eff", kind="prop", num="0.00")
# G45 = σ_cr,E 헬퍼 (탄성 임계응력)
ws.cell(45, 7, '=IF(AND($C$16>0,$C$7>0,$C$16*$C$13<$C$8),'
               'PI()^2*$C$23/(12*(1-$C$18^2))*($C$41/($C$11*$C$10))^2,"")')
ws.cell(45, 7).font = f_note
ws.cell(45, 6, "σ_cr,E 헬퍼 →").font = f_note
put(45, "압축 σ_cr (Johnson)",
    '=IF(ISNUMBER($G$45),IF($G$45<=$C$22/2,$G$45,'
    '$C$22*(1-$C$22/(4*$G$45))),"")', "MPa", "비탄성 천이", kind="prop",
    num="0.00")
put(46, "압축 허용 min(σ_cr/DF,S)",
    '=IF(ISNUMBER($C$45),MIN($C$45/$C$19,$C$21),"")', "MPa", "", kind="prop",
    num="0.00")
put(47, "  압축 판정",
    '=IF(AND($C$16>0,$C$7>0),IF(NOT(ISNUMBER($C$44)),"—",'
    'IF($C$44<=$C$46,"합격 ✓","불합격 ✗")),"")', "-", "", kind="res",
    num="General")
put(48, "■ 종합 판정",
    '=IF($C$5>204,"온도초과",IF($C$16>0,IF(OR($C$43="불합격 ✗",$C$47="불합격 ✗"),'
    '"■ 종합 불합격","■ 종합 합격"),"(선정두께 미입력 — 검증 생략)"))', "-",
    "", kind="res", num="General")
ws.cell(48, 3).font = Font(name=ARIAL, size=11, bold=True, color="1F4E78")

for r in list(range(5, 17)) + list(range(21, 33)) + list(range(34, 40)) \
        + list(range(41, 49)):
    ws.row_dimensions[r].height = 15

# ---------------- 강도계산서 (요약 리포트) ----------------
# 모든 값은 [계산기] 시트를 참조하는 수식 — 입력을 바꾸면 계산서도 자동 갱신된다.
ws_r = wb.create_sheet("강도계산서")
wb.move_sheet("강도계산서", -(len(wb.sheetnames) - 2))   # 계산기 바로 뒤
ws_r.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGH", (1.8, 4.5, 27, 27, 12, 8, 17, 17)):
    ws_r.column_dimensions[col].width = w
ws_r.page_setup.orientation = "portrait"
ws_r.page_setup.paperSize = ws_r.PAPERSIZE_A4
ws_r.sheet_properties.pageSetUpPr.fitToPage = True
ws_r.page_setup.fitToWidth = 1
ws_r.page_setup.fitToHeight = 0
ws_r.print_options.horizontalCentered = True
ws_r.page_margins.left = ws_r.page_margins.right = 0.4

f_rtitle = Font(name=ARIAL, size=15, bold=True, color="FFFFFF")
f_rsec = Font(name=ARIAL, size=10.5, bold=True, color="FFFFFF")
f_rh = Font(name=ARIAL, size=9, bold=True, color="FFFFFF")
f_rb = Font(name=ARIAL, size=9)
f_rbb = Font(name=ARIAL, size=9, bold=True)
f_rin = Font(name=ARIAL, size=9, bold=True, color="0000FF")
f_rsm = Font(name=ARIAL, size=8, color="444444")
fill_sec = PatternFill("solid", fgColor="1F4E78")
fill_th = PatternFill("solid", fgColor="4472C4")
fill_alt = PatternFill("solid", fgColor="F7F9FC")
med = Side(style="medium", color="1F4E78")

def rmerge(r, c1, c2, value=None):
    ws_r.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    cell = ws_r.cell(r, c1)
    if value is not None:
        cell.value = value
    return cell

def rsec(r, text):
    c = rmerge(r, 2, 8, text)
    c.font = f_rsec; c.fill = fill_sec
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_r.row_dimensions[r].height = 18

def rhead(r, labels, widths_cols):
    for txt, (c1, c2) in zip(labels, widths_cols):
        c = rmerge(r, c1, c2, txt) if c2 > c1 else ws_r.cell(r, c1, txt)
        c.font = f_rh; c.fill = fill_th; c.alignment = ctr; c.border = box
    ws_r.row_dimensions[r].height = 16

def rrow(r, cells, cols, *, alt=False, num=None, bold_col=None):
    """cells: 값/수식 리스트, cols: [(c1,c2), ...] 병합범위."""
    for i, (txt, (c1, c2)) in enumerate(zip(cells, cols)):
        c = rmerge(r, c1, c2, txt) if c2 > c1 else ws_r.cell(r, c1, txt)
        c.font = f_rbb if bold_col == i else f_rb
        c.border = box
        if alt:
            c.fill = fill_alt
        if num and num[i]:
            c.number_format = num[i]
            c.alignment = ctr if num[i] == "General" else right
        else:
            c.alignment = left if c1 in (3, 4) else ctr
    ws_r.row_dimensions[r].height = 15

# --- 표제 ---
t = rmerge(1, 2, 8, "강 도 계 산 서   /   STRENGTH CALCULATION SHEET")
t.font = f_rtitle; t.fill = fill_sec; t.alignment = ctr
ws_r.row_dimensions[1].height = 30
t2 = rmerge(2, 2, 8,
            "PFHE(Brazed Aluminium Plate-Fin Heat Exchanger) Plain 핀 두께 강도 검토 "
            "· ASME Sec. VIII Div.1 / ALPEMA Standard 2024")
t2.font = Font(name=ARIAL, size=8.5, italic=True, color="333333")
t2.alignment = ctr
ws_r.row_dimensions[2].height = 14

# --- 1. 문서 정보 ---
rsec(4, "1.  문 서 정 보  (Document Information)")
doc_rows = [
    ("프로젝트명", "", "문서번호", ""),
    ("기기번호 (Item No.)", "", "Rev.", "0"),
    ("작성자 (Prepared by)", "", "작성일 (Date)", ""),
    ("검토자 (Checked by)", "", "적용 코드", None),
]
for i, (l1, v1, l2, v2) in enumerate(doc_rows):
    r = 5 + i
    a = rmerge(r, 2, 3, l1); a.font = f_rbb; a.border = box; a.alignment = right
    b = rmerge(r, 4, 5, v1); b.font = f_rin; b.fill = fill_in; b.border = box
    b.alignment = left
    c = ws_r.cell(r, 6, l2); c.font = f_rbb; c.border = box; c.alignment = right
    d = rmerge(r, 7, 8)
    d.border = box; d.alignment = left
    if v2 is None:
        d.value = "ASME Sec. VIII Div.1 / ALPEMA 2024"
        d.font = f_rb
    else:
        d.value = v2; d.font = f_rin; d.fill = fill_in
    ws_r.row_dimensions[r].height = 15

# --- 2. 설계 조건 ---
rsec(10, "2.  설 계 조 건  (Design Conditions)")
COLS4 = [(2, 2), (3, 3), (4, 4), (5, 5), (6, 6), (7, 8)]
rhead(11, ["No.", "항 목", "기 호", "값", "단위", "비 고"], COLS4)
design = [
    ("설계온도", "T", "=계산기!$C$5", "°C", '=IF(계산기!$C$5>204,"⚠️ 재질한계 204°C 초과","≤ 204 (ALPEMA Table 6-1)")', "0.0"),
    ("핀 인장 유발 차압", "ΔP_t", "=계산기!$C$6", "MPa", "층 설계압력(MAWP), 인접층 0 가정", "0.000"),
    ("핀 압축 유발 차압", "ΔP_c", "=계산기!$C$7", "MPa", '=IF(계산기!$C$7>0,"인접층 가압 / 자층 감압 조건","해당 없음")', "0.000"),
    ("자층 핀 피치", "p", "=계산기!$C$8", "mm", "=계산기!$C$31", "0.000"),
    ("인접층 핀 피치", "p_adj", "=계산기!$C$9", "mm", '=IF(계산기!$C$9>0,"파팅시트 스팬 지배 검토","미입력 (자층 기준)")', "0.000"),
    ("핀 높이 (=사이드바 높이)", "h", "=계산기!$C$10", "mm", "=계산기!$C$32", "0.0"),
    ("좌굴 유효길이계수", "K", "=계산기!$C$11", "-", '=IF(계산기!$C$11<=0.7,"고정-힌지 준용 (sway 억제 전제)","보수적 적용")', "0.00"),
    ("인장 스테이 계수", "f", "=계산기!$C$12", "-", '=IF(계산기!$C$12>1,"UG-50(b) 준용 (+10%)","미적용")', "0.00"),
    ("제작 공차계수", "—", "=계산기!$C$13", "-", '=IF(계산기!$C$13<1,"유효두께 = 공칭 × 공차계수","공차 미반영")', "0.00"),
    ("파팅시트 층간 차압", "w", "=계산기!$C$14", "MPa", "|P_A − P_B|", "0.000"),
    ("균형가압 압력", "P_min", "=계산기!$C$15", "MPa", '=IF(계산기!$C$15>0,"핀 풋 어긋남 교번하중 반영","미고려")', "0.000"),
    ("선정 핀 두께 (공칭)", "t_sel", '=IF(계산기!$C$16>0,계산기!$C$16,"—")', "mm", '=IF(계산기!$C$16>0,"§5 검증 대상","미선정 — 검증 생략")', "0.000"),
]
for i, (nm, sym, val, unit, note, nf) in enumerate(design):
    r = 12 + i
    rrow(r, [i + 1, nm, sym, val, unit, note], COLS4, alt=(i % 2 == 1),
         num=["General", None, "General", nf, "General", None], bold_col=1)

# --- 3. 재료 및 허용응력 ---
rsec(25, "3.  재 료 및 허 용 응 력  (Material & Allowable Stress)")
rhead(26, ["No.", "항 목", "기 호", "값", "단위", "출 처 / 비 고"], COLS4)
mat = [
    ("핀 재질", "—", "AL 3003-O", "-", "SB-209, 브레이징 후 완전 어닐링(ALPEMA 5.13)", "General"),
    ("설계온도 허용응력", "S(T)", "=계산기!$C$21", "MPa", "ASME Sec. II-D Table 1B (온도 선형보간)", "0.00"),
    ("설계온도 항복강도", "Sy(T)", "=계산기!$C$22", "MPa", "Johnson 비탄성 좌굴 천이용", "0.00"),
    ("설계온도 탄성계수", "E(T)", "=계산기!$C$23", "MPa", "ASME Sec. II-D TM (좌굴 임계응력용)", "0"),
    ("포아송비", "ν", "=계산기!$C$18", "-", "알루미늄 합금", "0.00"),
    ("좌굴 설계계수", "DF", "=계산기!$C$19", "-", "ASME Div.1 UG-28 철학 준용", "0.0"),
]
for i, (nm, sym, val, unit, note, nf) in enumerate(mat):
    r = 27 + i
    rrow(r, [i + 1, nm, sym, val, unit, note], COLS4, alt=(i % 2 == 1),
         num=["General", None, "General", nf, "General", None], bold_col=1)

# --- 4. 강도 계산 결과 ---
rsec(34, "4.  강 도 계 산 결 과  (Calculation Results)")
COLS5 = [(2, 2), (3, 3), (4, 4), (5, 5), (6, 6), (7, 8)]
rhead(35, ["No.", "검 토 항 목", "적 용 식", "계산값", "단위",
           "기준 / 판정"], COLS5)
calcs = [
    ("핀 인장 최소두께", "t ≥ f·ΔP_t·p / (S + f·ΔP_t)", "=계산기!$C$26", "mm",
     '="식(2) — 지배: "&IF(계산기!$C$26>=계산기!$C$27,"■ 인장 지배","인장 비지배")'),
    ("핀 좌굴 최소두께", "ΔP_c·p/t ≤ min(σ_cr/DF, S)", "=계산기!$C$27", "mm",
     '=IF(계산기!$C$7<=0,"식(3~5) — ΔP_c=0, 해당 없음","식(3~5) Johnson 천이 — "&IF(계산기!$C$27>계산기!$C$26,"■ 좌굴 지배","좌굴 비지배"))'),
    ("요구 핀두께 (유효)", "max( 4.1 , 4.2 )", "=계산기!$C$28", "mm",
     "최소 보증두께 기준"),
    ("요구 핀두께 (공칭)", "유효두께 / 공차계수", "=계산기!$C$29", "mm",
     "=계산기!$C$30"),
    ("파팅시트 최소두께", "t_ps ≥ √( 4·M_d / S )", "=계산기!$C$38", "mm",
     "=계산기!$C$39"),
]
for i, (nm, eq, val, unit, judge) in enumerate(calcs):
    r = 36 + i
    rrow(r, [f"4.{i + 1}", nm, eq, val, unit, judge], COLS5, alt=(i % 2 == 1),
         num=["General", None, None, "0.0000", "General", None], bold_col=1)
# 파팅시트 보조 정보
rrow(41, ["", "  (파팅시트 중간값)",
          '="지배 피치 p̄ = "&TEXT(계산기!$C$34,"0.000")&" mm,  스팬 s = "&TEXT(계산기!$C$36,"0.000")&" mm"',
          "=계산기!$C$37", "N·mm/mm", "설계 모멘트 M_d = w·s²/12 + P_min·(p̄−t)·p̄/8"],
     COLS5, num=[None, None, None, "0.0000", "General", None])
for c in range(2, 9):
    ws_r.cell(41, c).font = f_rsm

# --- 5. 선정 사양 검증 ---
rsec(43, "5.  선 정 사 양 검 증  (Verification of Selected Fin)")
COLS6 = [(2, 2), (3, 3), (4, 4), (5, 5), (6, 6), (7, 7), (8, 8)]
rhead(44, ["No.", "검 토 항 목", "적 용 식", "계산 응력", "허용 응력",
           "안전율 SF", "판 정"], COLS6)
ver = [
    ("핀 인장 응력", "f·σ_t = f·ΔP_t·(p−t)/t",
     '=IF(계산기!$C$42="","—",계산기!$C$42)', "=계산기!$C$21",
     '=IF(AND(ISNUMBER(계산기!$C$42),계산기!$C$42>0),계산기!$C$21/계산기!$C$42,"—")',
     '=IF(계산기!$C$43="","—",계산기!$C$43)'),
    ("핀 압축·좌굴 응력", "σ_c = ΔP_c·p/t",
     '=IF(계산기!$C$44="","—",계산기!$C$44)',
     '=IF(계산기!$C$46="","—",계산기!$C$46)',
     '=IF(AND(ISNUMBER(계산기!$C$44),계산기!$C$44>0),계산기!$C$46/계산기!$C$44,"—")',
     '=IF(계산기!$C$47="","—",계산기!$C$47)'),
]
for i, (nm, eq, sig, allow, sf, judge) in enumerate(ver):
    r = 45 + i
    rrow(r, [f"5.{i + 1}", nm, eq, sig, allow, sf, judge], COLS6,
         alt=(i % 2 == 1),
         num=["General", None, None, "0.00", "0.00", "0.00", "General"],
         bold_col=1)
# 5.3 두께 여유
rrow(47, ["5.3", "핀 두께 여유", "유효 t / 요구 t (유효 기준)",
          '=IF(계산기!$C$16>0,계산기!$C$41,"—")',
          "=계산기!$C$28",
          '=IF(AND(계산기!$C$16>0,계산기!$C$28>0),계산기!$C$41/계산기!$C$28,"—")',
          '=IF(계산기!$C$16<=0,"—",IF(계산기!$C$41>=계산기!$C$28,"합격 ✓","불합격 ✗"))'],
     COLS6, num=["General", None, None, "0.000", "0.000", "0.00", "General"],
     bold_col=1)
# 5.4 좌굴 임계응력 참고
rrow(48, ["5.4", "  (참고) 좌굴 임계응력",
          "σ_cr : Johnson 포물선 천이 적용",
          '=IF(계산기!$C$45="","—",계산기!$C$45)', "—",
          '=IF(AND(ISNUMBER(계산기!$C$45),계산기!$C$44>0),계산기!$C$45/계산기!$C$44,"—")',
          "참고값"],
     COLS6, alt=True,
     num=["General", None, None, "0.00", "General", "0.00", "General"])
for c in range(2, 9):
    ws_r.cell(48, c).font = f_rsm

# --- 6. 결론 ---
rsec(50, "6.  결 론  (Conclusion)")
concl = [
    '="1)  설계조건 — 설계온도 "&TEXT(계산기!$C$5,"0")&" °C, 인장차압 ΔP_t = "'
    '&TEXT(계산기!$C$6,"0.000")&" MPa, 압축차압 ΔP_c = "&TEXT(계산기!$C$7,"0.000")'
    '&" MPa, 핀 피치 "&TEXT(계산기!$C$8,"0.000")&" mm, 핀 높이 "'
    '&TEXT(계산기!$C$10,"0.0")&" mm."',
    '="2)  허용응력 S(T) = "&TEXT(계산기!$C$21,"0.00")'
    '&" MPa (AL 3003-O, ASME Sec. II-D) 기준, 요구 핀두께는 "'
    '&IF(ISNUMBER(계산기!$C$29),TEXT(계산기!$C$29,"0.000")&" mm (공칭)",계산기!$C$29)'
    '&" 이며 ALPEMA 제작범위(0.15~0.7 mm) 대비 "&계산기!$C$30&" 이다."',
    '="3)  파팅시트 최소두께는 "&IF(ISNUMBER(계산기!$C$38),'
    'TEXT(계산기!$C$38,"0.000")&" mm",계산기!$C$38)'
    '&" 이며, ALPEMA 표준범위(0.8~2.0 mm) 대비 "&계산기!$C$39&" 이다."',
    '="4)  선정 핀두께 "&IF(계산기!$C$16>0,TEXT(계산기!$C$16,"0.000")&" mm (유효 "'
    '&TEXT(계산기!$C$41,"0.000")&" mm)","(미선정)")&" 에 대한 응력 검증 결과 — 인장 "'
    '&IF(계산기!$C$16>0,계산기!$C$43,"생략")&IF(계산기!$C$7>0,", 압축·좌굴 "'
    '&IF(계산기!$C$16>0,계산기!$C$47,"생략"),"")&" 로 평가되었다."',
]
for i, cf in enumerate(concl):
    r = 51 + i
    c = rmerge(r, 2, 8, cf)
    c.font = f_rb
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_r.row_dimensions[r].height = 15

# 최종 판정 박스
fin = rmerge(56, 2, 8,
             '="■  최 종 판 정  :   "&계산기!$C$48&'
             'IF(계산기!$C$5>204,"   (설계온도 재질한계 초과 — 결과 무효)","")')
fin.font = Font(name=ARIAL, size=12, bold=True, color="1F4E78")
fin.fill = PatternFill("solid", fgColor="DDEBF7")
fin.alignment = ctr
fin.border = Border(left=med, right=med, top=med, bottom=med)
ws_r.row_dimensions[56].height = 26

# --- 7. 적용 기준 및 한계 ---
rsec(58, "7.  적 용 기 준 및 한 계  (Basis & Limitations)")
limits = [
    "· 본 계산서는 예비설계용이다. 인증 설계는 ALPEMA 5.15.1.1(코드 당국 승인 계산법) "
    "또는 5.15.1.2(파열시험)로 확정하여야 한다.",
    "· 정적 압력하중 전용 — 열응력·피로는 ALPEMA 8.1의 운전 제한(±1 °C/min 등)으로 별도 관리한다.",
    "· 허용응력 S(T)·항복강도 Sy(T)·탄성계수 E(T)는 대표값이며, 적용 코드 연도판의 "
    "ASME Sec. II-D 값으로 검증·교체하여야 한다. ([물성DB] 시트)",
    "· 149 °C 이상 장기 운전은 크리프 영향으로 별도 검토가 필요하다. 204 °C 초과는 3003 재질 적용범위 밖이다.",
    "· 대상은 Plain 핀에 한한다. Serrated·Perforated 핀, 디스트리뷰터 핀, 헤더·노즐부는 범위 밖이다.",
    "· 인접층 압력은 신뢰할 수 있는 감압 방지 수단이 없는 한 0(대기압)으로 가정하였다 "
    "(ALPEMA 5.6.1 개별 챔버 가압 / 5.6.2.1 단일 챔버 설계압력 근거).",
    "· 상세 유도 근거: fin_thickness_derivation.md,  계산 로직 원본: fin_thickness_calc.py",
]
for i, txt in enumerate(limits):
    r = 59 + i
    c = rmerge(r, 2, 8, txt)
    c.font = f_rsm
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_r.row_dimensions[r].height = 13

# --- 서명란 ---
sign_r = 59 + len(limits) + 1
rhead(sign_r, ["", "작 성 (Prepared)", "검 토 (Checked)", "", "승 인 (Approved)",
               "", "일 자 (Date)"],
      [(2, 2), (3, 3), (4, 5), (6, 6), (7, 7), (8, 8), (8, 8)])
for c1, c2 in [(3, 3), (4, 5), (7, 8)]:
    cell = rmerge(sign_r + 1, c1, c2)
    cell.border = box
ws_r.cell(sign_r + 1, 2).border = box
ws_r.cell(sign_r + 1, 6).border = box
ws_r.row_dimensions[sign_r + 1].height = 34
foot = rmerge(sign_r + 3, 2, 8,
              '="본 계산서는 [계산기] 시트의 입력값에 연동되어 자동 산출됨.   '
              '적용 코드: ASME Sec. VIII Div.1 · ALPEMA Standard 2024   |   '
              '재질: AL 3003-O   |   단위: MPa, mm"')
foot.font = f_rsm
foot.alignment = ctr

# ---------------- 설명 시트 ----------------
ws_n = wb.create_sheet("설명")
ws_n.sheet_view.showGridLines = False
ws_n.column_dimensions["A"].width = 100
notes = [
    ("PFHE Plain 핀 적정두께 계산기 — 사용 안내", 12, True, "1F4E78"),
    ("", 10, False, None),
    ("■ 사용법", 11, True, "2E75B6"),
    ("1. [계산기] 시트의 파랑(노랑배경) 셀만 편집합니다. 나머지는 자동 계산 수식입니다.", 10, False, None),
    ("2. 단위는 압력 MPa, 길이 mm 입니다 (1 MPa = 10 bar ≈ 10.2 kgf/cm²).", 10, False, None),
    ("3. [1][2]는 최소 소요두께, [3]은 채택 기준(둘 중 큰 값)입니다.", 10, False, None),
    ("4. t_sel(선정 두께)을 입력하면 ⑤ 검증에서 합/부를 판정합니다. 0이면 검증을 건너뜁니다.", 10, False, None),
    ("5. ΔP_c(압축)가 0이면 좌굴 검토는 생략됩니다.", 10, False, None),
    ("", 10, False, None),
    ("■ 계산 근거 (수식)", 11, True, "2E75B6"),
    ("[1] 핀 인장:    t ≥ f·ΔP_t·p / (S(T) + f·ΔP_t)", 10, False, None),
    ("[2] 핀 좌굴:    ΔP_c·p/t ≤ min(σ_cr/DF, S), σ_cr는 오일러식 + Johnson 비탄성 천이", 10, False, None),
    ("               → [좌굴스캔] 시트가 후보 두께를 훑어 조건을 만족하는 최소 t를 MINIFS로 반환", 10, False, None),
    ("[3] 파팅시트:  t_ps ≥ √(4·M_d/S),  M_d = w·s²/12 + P_min·(p̄−t)·p̄/8", 10, False, None),
    ("온도 반영:     S(T)/Sy(T)/E(T)를 [물성DB] 표에서 온도 선형보간 (204°C 초과 무효)", 10, False, None),
    ("", 10, False, None),
    ("■ 시트 구성", 11, True, "2E75B6"),
    ("계산기     — 입력/결과 (메인). 여기서만 입력합니다.", 10, False, None),
    ("강도계산서 — 계산 결과를 문서 형식으로 요약 (설계조건·재료·계산결과·검증·결론·서명란).", 10, False, None),
    ("             [계산기] 입력에 자동 연동되며, A4 세로 1페이지 폭에 맞춰 인쇄됩니다.", 10, False, None),
    ("             문서정보(프로젝트명·기기번호·작성자 등) 파랑 셀은 직접 입력하세요.", 10, False, None),
    ("물성DB     — S(T)/Sy(T)/E(T) 표 (코드 연도판으로 교체 대상)", 10, False, None),
    ("좌굴스캔   — 좌굴 최소두께 반복해용 스캔 테이블 (수정 불필요)", 10, False, None),
    ("", 10, False, None),
    ("■ 한계 및 주의", 11, True, "C00000"),
    ("· 예비설계용. 정적 압력하중 전용(열응력·피로는 ALPEMA 8.1 운전제한으로 별도 관리).", 10, False, None),
    ("· S/Sy/E 표는 대표값 — 적용 코드 연도판으로 반드시 검증/교체할 것.", 10, False, None),
    ("· 149°C 이상 장기운전은 크리프 별도 검토. Serrated/천공핀·헤더·노즐은 범위 밖.", 10, False, None),
    ("· 인증 설계는 ALPEMA 5.15.1.1(승인 계산법)/5.15.1.2(파열시험)으로 확정.", 10, False, None),
    ("· 상세 유도는 fin_thickness_derivation.md, 원본 로직은 fin_thickness_calc.py 참조.", 10, False, None),
]
for i, (txt, sz, bold, color) in enumerate(notes, 1):
    c = ws_n.cell(i, 1, txt)
    c.font = Font(name=ARIAL, size=sz, bold=bold, color=color or "000000")
    c.alignment = left

# 열 때 전체 재계산 강제 → 어떤 뷰어에서도 값이 즉시 표시됨
wb.calculation.fullCalcOnLoad = True

wb.save(OUT)
print(f"saved: {OUT}")
print(f"scan rows: {n}")
